"""Convert seed/bank_soal.xlsx (sheet "Bank Soal") into seed/questions.json.

Usage: python seed/convert_bank.py   (needs: pip install openpyxl)
The xlsx is the editable source (it also holds sumber / status_validasi, which are not seeded); questions.json is what `make seed` loads.
"""
import json
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).parent
LETTERS = "ABCD"


def clean(value):
    return "" if value is None else str(value).strip()


def build_options(row):
    options = []
    for i, letter in enumerate(LETTERS):
        label_id, label_en = clean(row[f"opsi_{letter.lower()}_id"]), clean(row[f"opsi_{letter.lower()}_en"])
        if not label_id and not label_en:
            continue
        options.append({
            "id": f"opt-{letter.lower()}",
            "label": {"id": label_id, "en": label_en},
            "correct": letter == clean(row["jawaban_benar"]).upper(),
        })
    return options


def convert_row(row):
    return {
        "id": clean(row["id"]),
        "site": int(row["site"]),
        "level": int(row["level"]),
        "type": clean(row["type"]),
        "prompt": {"id": clean(row["prompt_id"]), "en": clean(row["prompt_en"])},
        "options": build_options(row),
        "explanation": {"id": clean(row["penjelasan_id"]), "en": clean(row["penjelasan_en"])},
        "active": clean(row["active"]).upper() == "TRUE",
    }


def main():
    ws = openpyxl.load_workbook(HERE / "bank_soal.xlsx", data_only=True)["Bank Soal"]
    header = [c.value for c in ws[1]]
    rows = [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
    questions = [convert_row(r) for r in rows]
    out = HERE / "questions.json"
    out.write_text(json.dumps(questions, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"wrote {len(questions)} questions to {out}", file=sys.stderr)


if __name__ == "__main__":
    main()
